"""Test entity linking (offline) - synthetic records, no network."""
import csv as _csv
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

from models import IdentityRecord
from scraper.entity_linker import (
    canonical_fb,
    export_entities_csv,
    export_entities_workbook,
    link_identities,
)


def _records():
    return [
        IdentityRecord(key="a", name="Nguyen Van A", profile_url="facebook.com/nguyena",
                       phones=["+84912345678"], socials=["facebook|nguyena"],
                       group_ids=["111"], verified=True),
        IdentityRecord(key="b", name="Nguyen Va", profile_url="facebook.com/profile.php?id=100023456789",
                       phones=["+84912345678"], emails=["shared@gmail.com"],
                       group_ids=["111", "222"]),
        IdentityRecord(key="c", name="N. Van A", profile_url="facebook.com/nguyena2",
                       emails=["shared@gmail.com"], socials=["instagram|nguyena"],
                       group_ids=["333"]),
        IdentityRecord(key="d", name="Tran Thi D", profile_url="facebook.com/trangd",
                       socials=["tiktok|trand.official"], group_ids=["111"]),
        IdentityRecord(key="e", name="Tran D", profile_url="facebook.com/trand2",
                       socials=["tiktok|@trand.official"], group_ids=["444"]),
        IdentityRecord(key="f", name="Le F", profile_url="facebook.com/lef",
                       socials=["instagram|lef.ig"]),
    ]


def test_canonical_fb():
    assert canonical_fb("https://www.facebook.com/nguyena?ref=x") == "nguyena"
    assert canonical_fb("https://m.facebook.com/NguyenA/") == "nguyena"
    assert canonical_fb("https://www.facebook.com/profile.php?id=100023456789") == "100023456789"
    assert canonical_fb("https://www.facebook.com/groups/123456789") is None
    assert canonical_fb(None) is None


def test_link_identities_clusters():
    clusters = link_identities(_records())
    sizes = sorted(len(c.members) for c in clusters)
    assert sizes == [2, 3], sizes
    big = max(clusters, key=lambda c: len(c.members))
    assert big.confidence == "certain"
    assert {m.key for m in big.members} == {"a", "b", "c"}
    signals = " ".join(big.shared_signals)
    assert "+84912345678" in signals
    assert "shared@gmail.com" in signals


def test_link_identities_social_handle_dedupe():
    # handle trùng (có @ và không @) phải nhận ra
    clusters = link_identities(_records())
    small = min(clusters, key=lambda c: len(c.members))
    assert small.confidence in {"certain", "high"}
    assert {m.key for m in small.members} == {"d", "e"}


def test_no_false_positive():
    recs = [
        IdentityRecord(key="x", name="X", profile_url="facebook.com/x", phones=["+84111111111"]),
        IdentityRecord(key="y", name="Y", profile_url="facebook.com/y", phones=["+84222222222"]),
    ]
    assert link_identities(recs) == []


def test_export(tmp_path):
    clusters = link_identities(_records())
    csv_path = export_entities_csv(clusters, tmp_path / "clusters.csv")
    xlsx_path = export_entities_workbook(clusters, tmp_path / "clusters.xlsx")
    with csv_path.open(encoding="utf-8-sig") as f:
        rows = list(_csv.DictReader(f))
    assert len(rows) == 2, rows
    assert rows[0]["Số tài khoản"] in {"2", "3"}
    assert "Mã cụm" in rows[0] and "Tín hiệu liên kết" in rows[0]
    assert xlsx_path.exists() and xlsx_path.stat().st_size > 1000

    import openpyxl
    wb = openpyxl.load_workbook(xlsx_path)
    assert wb.sheetnames == ["Cum danh tinh"], wb.sheetnames
    ws = wb["Cum danh tinh"]
    header = [c.value for c in ws[1]]
    assert "Tên hiển thị" in header and "SĐT liên kết" in header, header


if __name__ == "__main__":
    import pytest
    sys.exit(pytest.main([__file__, "-v"]))
